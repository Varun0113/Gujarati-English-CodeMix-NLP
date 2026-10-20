"""
Complaint_Intelligence_NLP_Project Generator
Generates a 20-slide, 16:9 widescreen, professional, academic MCA presentation.
"""

import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Semantic Color Palette
BG_COLOR = RGBColor(247, 247, 245)      # #F7F7F5
SURFACE = RGBColor(255, 255, 255)       # #FFFFFF
BORDER = RGBColor(230, 232, 236)        # #E6E8EC
PRIMARY = RGBColor(23, 32, 51)          # #172033 (Deep Navy Slate)
SECONDARY = RGBColor(102, 112, 133)     # #667085 (Slate Gray)
INDIGO = RGBColor(79, 70, 229)          # #4F46E5
DEEP_INDIGO = RGBColor(55, 48, 163)     # #3730A3
SOFT_INDIGO = RGBColor(238, 242, 255)   # #EEF2FF
TEAL = RGBColor(15, 118, 110)           # #0F766E
SOFT_TEAL = RGBColor(236, 253, 245)     # #ECFDF5
POSITIVE = RGBColor(22, 128, 60)        # #16803C
NEGATIVE = RGBColor(194, 65, 59)        # #C2413B
WARNING = RGBColor(183, 121, 31)        # #B7791F
SOFT_AMBER = RGBColor(254, 243, 199)    # #FEF3C7
SOFT_RED = RGBColor(254, 226, 226)      # #FEE2E2
DARK_NAVY = RGBColor(15, 23, 42)        # #0F172A

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"

def set_slide_background(slide, color=BG_COLOR):
    """Draw full-slide background rectangle."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, category_tag, title_text, subtitle_text, slide_num, total_slides=20):
    """Add standard academic header spanning full slide width."""
    # Top Category Pill
    tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.6), Inches(0.32))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = SOFT_INDIGO
    tag_box.line.color.rgb = INDIGO
    tag_box.line.width = Pt(1)
    tf = tag_box.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = category_tag.upper()
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = INDIGO
    p.alignment = PP_ALIGN.CENTER

    # Slide Number Badge (Top Right)
    num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.433), Inches(0.4), Inches(1.1), Inches(0.32))
    num_box.fill.solid()
    num_box.fill.fore_color.rgb = SURFACE
    num_box.line.color.rgb = BORDER
    num_box.line.width = Pt(1)
    tf_num = num_box.text_frame
    tf_num.word_wrap = False
    tf_num.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_num = tf_num.paragraphs[0]
    p_num.text = f"{slide_num:02d} / {total_slides:02d}"
    p_num.font.name = FONT_HEADING
    p_num.font.size = Pt(10)
    p_num.font.bold = True
    p_num.font.color.rgb = SECONDARY
    p_num.alignment = PP_ALIGN.CENTER

    # Title & Subtitle Box
    header_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.75), Inches(11.783), Inches(0.8))
    tf_head = header_box.text_frame
    tf_head.word_wrap = True
    tf_head.margin_left = tf_head.margin_right = tf_head.margin_top = tf_head.margin_bottom = 0

    p_title = tf_head.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(21)
    p_title.font.bold = True
    p_title.font.color.rgb = PRIMARY
    p_title.space_after = Pt(2)

    p_sub = tf_head.add_paragraph()
    p_sub.text = subtitle_text
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = SECONDARY

    # Horizontal Divider Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.58), Inches(11.733), Pt(1))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER
    line.line.fill.background()

def add_card(slide, left, top, width, height, bg_color=SURFACE, border_color=BORDER, border_width=1):
    """Draw a card panel."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_badge(slide, left, top, width, height, text, bg_color, text_color, font_size=10, bold=True):
    """Small pill badge."""
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEADING
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return badge

def add_footer_note(slide, note_text):
    """Add a professional academic footer note."""
    box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = note_text
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.color.rgb = SECONDARY

print("Base setup ready.")
